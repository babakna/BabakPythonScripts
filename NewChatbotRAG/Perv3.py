import sys
import os
import subprocess
import PyPDF2
import docx
import openpyxl
import requests
import json
import chromadb
import psutil
import time
from PyQt5.QtWidgets import QApplication, QWidget, QPushButton, QVBoxLayout, QHBoxLayout, QTextEdit, QFileDialog, QComboBox, QLabel, QMessageBox
from PyQt5.QtCore import QThread, pyqtSignal

class OllamaThread(QThread):
    finished = pyqtSignal(bool, str)

    def run(self):
        try:
            # Check if Ollama is already running
            for proc in psutil.process_iter(['name']):
                if proc.info['name'] == 'ollama':
                    self.finished.emit(True, "Ollama is already running.")
                    return

            # Start Ollama
            subprocess.Popen(["ollama", "serve"], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
            
            # Wait for Ollama to start (max 30 seconds)
            for _ in range(30):
                try:
                    requests.get("http://localhost:11434/api/tags")
                    self.finished.emit(True, "Ollama started successfully.")
                    return
                except requests.RequestException:
                    time.sleep(1)
            
            self.finished.emit(False, "Failed to start Ollama after 30 seconds.")
        except Exception as e:
            self.finished.emit(False, f"Error starting Ollama: {str(e)}")

class AIAssistant(QWidget):
    def __init__(self):
        super().__init__()
        self.initUI()
        self.documents = {}
        self.chroma_client = chromadb.Client()
        self.collection = self.chroma_client.create_collection(name="documents")
        self.start_ollama()

    def initUI(self):
        self.setWindowTitle('AI Assistant')
        self.setGeometry(100, 100, 600, 400)

        layout = QVBoxLayout()

        # LLM selection
        self.llm_combo = QComboBox()
        self.llm_combo.addItem("llama2")  # Default option
        layout.addWidget(QLabel("Select LLM:"))
        layout.addWidget(self.llm_combo)

        # Buttons
        button_layout = QHBoxLayout()
        self.load_btn = QPushButton('Load Documents')
        self.display_btn = QPushButton('Display Documents')
        self.test_btn = QPushButton('Test Ollama')
        button_layout.addWidget(self.load_btn)
        button_layout.addWidget(self.display_btn)
        button_layout.addWidget(self.test_btn)
        layout.addLayout(button_layout)

        # Query input
        self.query_input = QTextEdit()
        self.query_input.setPlaceholderText("Enter your query here...")
        layout.addWidget(self.query_input)

        # Run query button
        self.run_btn = QPushButton('Run Query')
        layout.addWidget(self.run_btn)

        # Display area
        self.display = QTextEdit()
        self.display.setReadOnly(True)
        layout.addWidget(self.display)

        # Clear and Exit buttons
        bottom_layout = QHBoxLayout()
        self.clear_btn = QPushButton('Clear Display')
        self.exit_btn = QPushButton('Exit')
        bottom_layout.addWidget(self.clear_btn)
        bottom_layout.addWidget(self.exit_btn)
        layout.addLayout(bottom_layout)

        self.setLayout(layout)

        # Connect buttons to functions
        self.load_btn.clicked.connect(self.load_documents)
        self.display_btn.clicked.connect(self.display_documents)
        self.test_btn.clicked.connect(self.test_ollama)
        self.run_btn.clicked.connect(self.run_query)
        self.clear_btn.clicked.connect(self.clear_display)
        self.exit_btn.clicked.connect(self.exit_app)

    def start_ollama(self):
        self.ollama_thread = OllamaThread()
        self.ollama_thread.finished.connect(self.on_ollama_started)
        self.ollama_thread.start()

    def on_ollama_started(self, success, message):
        if success:
            self.display.append(message)
            self.populate_llm_list()
        else:
            QMessageBox.critical(self, "Error", message)

    def populate_llm_list(self):
        try:
            response = requests.get("http://localhost:11434/api/tags")
            if response.status_code == 200:
                models = response.json()['models']
                self.llm_combo.clear()
                for model in models:
                    self.llm_combo.addItem(model['name'])
            else:
                self.display.append(f"Error fetching LLM list: {response.status_code}")
        except requests.RequestException as e:
            self.display.append(f"Error connecting to Ollama: {str(e)}")

    def load_documents(self):
        files, _ = QFileDialog.getOpenFileNames(self, "Select Documents", "", "Documents (*.pdf *.docx *.xlsx)")
        for file in files:
            filename = os.path.basename(file)
            content = self.read_file(file)
            self.documents[filename] = content
            self.collection.add(
                documents=[content],
                metadatas=[{"source": filename}],
                ids=[filename]
            )
        self.display.append(f"Loaded {len(files)} documents.")

    def read_file(self, file_path):
        if file_path.endswith('.pdf'):
            return self.read_pdf(file_path)
        elif file_path.endswith('.docx'):
            return self.read_docx(file_path)
        elif file_path.endswith('.xlsx'):
            return self.read_excel(file_path)

    def read_pdf(self, file_path):
        with open(file_path, 'rb') as file:
            reader = PyPDF2.PdfReader(file)
            return " ".join(page.extract_text() for page in reader.pages)

    def read_docx(self, file_path):
        doc = docx.Document(file_path)
        return " ".join(para.text for para in doc.paragraphs)

    def read_excel(self, file_path):
        workbook = openpyxl.load_workbook(file_path)
        return " ".join(" ".join(str(cell.value) for cell in row if cell.value) for sheet in workbook for row in sheet.iter_rows())

    def display_documents(self):
        self.display.clear()
        for filename, content in self.documents.items():
            self.display.append(f"Document: {filename}")
            self.display.append(content[:500] + "...\n")

    def test_ollama(self):
        selected_model = self.llm_combo.currentText()
        response = self.query_ollama("Hello, are you working?", selected_model)
        self.display.append(f"Ollama Test Response ({selected_model}):\n{response}")

    def run_query(self):
        query = self.query_input.toPlainText()
        if not query:
            self.display.append("Please enter a query.")
            return

        # Retrieve relevant documents from ChromaDB
        results = self.collection.query(query_texts=[query], n_results=3)
        context = "Relevant information from documents:\n"
        for doc, metadata in zip(results['documents'][0], results['metadatas'][0]):
            context += f"From {metadata['source']}:\n{doc[:500]}...\n\n"

        # Prepare the prompt for Ollama
        prompt = f"{context}\nBased on the above information and your knowledge, please answer the following question:\n{query}"

        # Query Ollama
        selected_model = self.llm_combo.currentText()
        response = self.query_ollama(prompt, selected_model)

        self.display.append(f"Query: {query}")
        self.display.append(f"AI Response:\n{response}")

    def query_ollama(self, prompt, model):
        url = "http://localhost:11434/api/generate"
        data = {
            "model": model,
            "prompt": prompt
        }
        try:
            response = requests.post(url, json=data)
            if response.status_code == 200:
                return json.loads(response.text)['response']
            else:
                return f"Error: {response.status_code}"
        except requests.RequestException as e:
            return f"Error connecting to Ollama: {str(e)}"
        except json.JSONDecodeError:
            return "Error: Invalid response from Ollama"

    def clear_display(self):
        self.display.clear()

    def exit_app(self):
        self.chroma_client.delete_collection("documents")
        QApplication.quit()

if __name__ == '__main__':
    app = QApplication(sys.argv)
    ex = AIAssistant()
    ex.show()
    sys.exit(app.exec_())
